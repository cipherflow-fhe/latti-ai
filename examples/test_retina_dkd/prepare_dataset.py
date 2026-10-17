"""Select ten documented test samples and run the supplied ONNX segmenters."""

import argparse
import csv
import hashlib
import importlib.util
import json
from pathlib import Path
import shutil
import time

import numpy as np
import openpyxl


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        '--client-dir', type=Path, default=Path('/LattiAI-Task-Resources/retina_degree8_02/task/client')
    )
    parser.add_argument('--xlsx', type=Path, default=Path('/home/wangy/Retina-DKD/data/risk_factor_5.xlsx'))
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--cpu-count', type=int, default=4)
    args = parser.parse_args()
    client = args.client_dir.resolve()
    output = args.output_dir or client / 'demo_10'
    source = client / 'py_output_all_ub1p75_deg8_ep005.csv'
    reference = load_module('retina_preprocessing', client / 'gen_csv/gen_client_inputs_seg.py')
    segmentation = load_module('retina_segmentation', client / 'seg_inference_deploy/predict_onnx.py')
    with source.open(newline='', encoding='utf-8-sig') as file:
        rows = list(csv.DictReader(file))
    ids, non_inv, invasive = reference.read_xlsx_5(str(args.xlsx))
    selected, patients = [], set()
    # Largest absolute logit margin first; one image per patient, balanced classes.
    for label, correct, count in ((0, True, 4), (0, False, 1), (1, True, 5)):
        candidates = sorted(
            (
                row
                for row in rows
                if int(row['ground_truth']) == label and (row['ground_truth'] == row['predicted']) == correct
            ),
            key=lambda row: (-abs(float(row['raw_output_1']) - float(row['raw_output_0'])), row['image_path']),
        )
        group = []
        for row in candidates:
            image = Path(row['image_path'])
            pid = reference.get_id(image.name)
            if pid in patients:
                continue
            index = ids[pid]
            clinical = non_inv[index] + invasive[index]
            expected = [float(row[key]) for key in ('non_inv_0', 'non_inv_1', 'inv_0', 'inv_1', 'inv_2')]
            np.testing.assert_allclose(clinical, expected, rtol=0, atol=1e-15)
            assert int(np.argmax([float(row['raw_output_0']), float(row['raw_output_1'])])) == int(row['predicted'])
            patients.add(pid)
            group.append(row)
            if len(group) == count:
                break
        if len(group) != count:
            raise ValueError(f'Not enough distinct patients for label={label}, correct={correct}')
        selected.extend(group)
    selected.sort(key=lambda row: Path(row['image_path']).name)
    assert len(selected) == 10 and sum(row['ground_truth'] == row['predicted'] for row in selected) == 9

    for folder in ('images', 'disk', 'lesion'):
        (output / 'dataset' / folder).mkdir(parents=True, exist_ok=True)
    (output / 'reference_logits').mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    original_rows = list(workbook.worksheets[0].values)
    subset = openpyxl.Workbook()
    clinical_sheet = subset.active
    clinical_sheet.title = 'clinical_factors'
    clinical_sheet.append(original_rows[0][:8])
    for row in original_rows[1:]:
        patient = row[1] if isinstance(row[1], str) else str(int(row[1]))
        if patient in patients:
            clinical_sheet.append(row[:8])
    assert clinical_sheet.max_row == 11
    subset.save(output / 'dataset/risk_factor_5.xlsx')
    workbook.close()
    with (output / 'dataset/label.csv').open('w', newline='', encoding='utf-8') as file:
        writer = csv.writer(file)
        writer.writerow(['filename', 'label'])
        writer.writerows((Path(row['image_path']).name, row['ground_truth']) for row in selected)
    with (output / 'selected_samples.csv').open('w', newline='', encoding='utf-8') as file:
        writer = csv.DictWriter(file, fieldnames=['filename', 'patient_id'] + list(rows[0]))
        writer.writeheader()
        writer.writerows(
            dict(
                filename=Path(row['image_path']).name, patient_id=reference.get_id(Path(row['image_path']).name), **row
            )
            for row in selected
        )

    onnx_dir = client / 'seg_inference_deploy/onnx'
    providers = ['CPUExecutionProvider']
    lesion_session = segmentation.make_session(onnx_dir / 'lesion.onnx', providers, args.cpu_count)
    disc_session = segmentation.make_session(onnx_dir / 'disc.onnx', providers, args.cpu_count)
    cup_sessions = [
        segmentation.make_session(onnx_dir / f'cup{i}_b1.onnx', providers, args.cpu_count) for i in range(1, 6)
    ]
    hashes = {}
    for index, row in enumerate(selected, 1):
        start = time.monotonic()
        image = Path(row['image_path'])
        copied = output / 'dataset/images' / image.name
        shutil.copyfile(image, copied)
        assert copied.read_bytes() == image.read_bytes()
        hashes[image.name] = hashlib.sha256(image.read_bytes()).hexdigest()
        lesion = segmentation.predict_lesion(copied, lesion_session)
        segmentation.write_gray(output / 'dataset/lesion' / f'{image.stem}_lesion.jpg', lesion)
        disk = segmentation.predict_disk(copied, disc_session, cup_sessions, use_tta=False)
        segmentation.write_gray(output / 'dataset/disk' / f'{image.stem}_disk.jpg', disk)
        (output / 'reference_logits' / f'{image.stem}.csv').write_text(
            f'{row["raw_output_0"]},{row["raw_output_1"]}\n', encoding='utf-8'
        )
        print(f'[{index}/10] {image.name}: ONNX lesion + disk, {time.monotonic() - start:.2f}s', flush=True)

    manifest = {
        'source_predictions': str(source),
        'source_workbook': str(args.xlsx),
        'selection': 'Distinct patients; 5 DN + 5 NDRD; 9 correct + 1 incorrect; descending absolute logit margin within each group',
        'reference_total': 10,
        'reference_correct': 9,
        'reference_accuracy': 0.9,
        'accuracy_basis': 'Provided historical raw_output_0/raw_output_1, not newly executed classification inference',
        'labels': {'0': 'DN', '1': 'NDRD (including DN combined with NDRD)'},
        'segmentation': {
            'entry': str(client / 'seg_inference_deploy/predict_onnx.py'),
            'providers': providers,
            'cpu_count': args.cpu_count,
            'lesion_size': 1024,
            'binary_lesion': False,
            'disk_threshold': 0.75,
            'fast_disk': True,
            'models_sha256': {
                path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(onnx_dir.glob('*.onnx'))
            },
        },
        'original_images_sha256': hashes,
    }
    (output / 'manifest.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(f'Dataset ready: {output}', flush=True)


if __name__ == '__main__':
    main()
